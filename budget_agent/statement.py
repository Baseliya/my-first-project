"""Читання виписок і таблиць витрат (CSV / XLSX) у текстові фрагменти для Claude."""

import csv
import io

from openpyxl import load_workbook

CHUNK_ROWS = 120


def _decode(data: bytes) -> str:
    for enc in ("utf-8-sig", "cp1251"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def read_rows(data: bytes, filename: str) -> list[list[str]]:
    name = filename.lower()
    if name.endswith((".xlsx", ".xlsm")):
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        rows = []
        for ws in wb.worksheets:
            for row in ws.iter_rows(values_only=True):
                cells = ["" if v is None else str(v).strip() for v in row]
                if any(cells):
                    rows.append(cells)
        return rows
    text = _decode(data)
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    return [[c.strip() for c in row] for row in csv.reader(io.StringIO(text), dialect) if any(c.strip() for c in row)]


def to_chunks(rows: list[list[str]], chunk_rows: int = CHUNK_ROWS) -> list[str]:
    """Розбиває таблицю на частини; рядок заголовка повторюється в кожній частині."""
    if not rows:
        return []
    header, body = rows[0], rows[1:] or []
    if not body:
        return [" | ".join(header)]
    chunks = []
    for i in range(0, len(body), chunk_rows):
        part = [header] + body[i : i + chunk_rows]
        chunks.append("\n".join(" | ".join(r) for r in part))
    return chunks
